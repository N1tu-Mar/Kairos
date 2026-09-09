"""Bounded extraction for untrusted founder intake documents.

The original upload exists only in a temporary file owned by this module.
Parsing runs in a disposable subprocess with a deadline so a malformed PDF
or Office archive cannot pin an API worker indefinitely. Only sanitized,
bounded text chunks are returned to persistence.
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any

from pypdf import PdfReader
from pptx import Presentation

from agent.models import IntakeDocumentChunk
from agent.sanitize import ingest

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_DECOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 2_000
MAX_COMPRESSION_RATIO = 100
MAX_PDF_PAGES = 100
MAX_PPTX_SLIDES = 100
MAX_EXTRACTED_CHARS = 200_000
MAX_CHUNKS = 100
PARSER_TIMEOUT_SECONDS = 12

_ALLOWED_MEDIA_TYPES: dict[str, frozenset[str]] = {
    ".pdf": frozenset({"application/pdf"}),
    ".pptx": frozenset(
        {
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        }
    ),
    ".txt": frozenset({"text/plain"}),
    ".md": frozenset({"text/markdown", "text/plain"}),
    ".markdown": frozenset({"text/markdown", "text/plain"}),
}
_EXECUTABLE_SUFFIXES = {
    ".bat", ".cmd", ".com", ".dll", ".exe", ".hta", ".jar", ".js",
    ".lnk", ".msi", ".ps1", ".scr", ".sh", ".vbs",
}


class DocumentRejected(ValueError):
    """A safe, founder-facing reason an upload was refused."""


def safe_filename(filename: str) -> str:
    """Return a display-only basename without paths or control characters."""
    normalized = filename.replace("\\", "/")
    name = PurePosixPath(normalized).name.strip()
    if not name or name in {".", ".."} or any(ord(ch) < 32 for ch in name):
        raise DocumentRejected("the file name is invalid")
    if len(name) > 200:
        raise DocumentRejected("the file name is too long")
    return name


def _validate_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(_BytesReader(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise DocumentRejected("the presentation contains too many entries")
            names = {info.filename for info in infos}
            if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
                raise DocumentRejected("the file is not a valid PPTX presentation")
            expanded = 0
            for info in infos:
                path = PurePosixPath(info.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise DocumentRejected("the presentation contains an unsafe path")
                if info.flag_bits & 0x1:
                    raise DocumentRejected("encrypted presentations are not supported")
                suffix = path.suffix.casefold()
                if suffix in _EXECUTABLE_SUFFIXES or path.name.casefold() == "vbaproject.bin":
                    raise DocumentRejected("macro-enabled or executable files are not supported")
                expanded += info.file_size
                if expanded > MAX_DECOMPRESSED_BYTES:
                    raise DocumentRejected("the presentation expands beyond the safe limit")
                if info.file_size and info.compress_size == 0:
                    raise DocumentRejected("the presentation has an unsafe compression ratio")
                if info.compress_size and info.file_size / info.compress_size > MAX_COMPRESSION_RATIO:
                    raise DocumentRejected("the presentation has an unsafe compression ratio")
    except (zipfile.BadZipFile, OSError):
        raise DocumentRejected("the file is not a valid PPTX presentation") from None


class _BytesReader:
    """Small seekable wrapper accepted by ZipFile without another full copy."""

    def __init__(self, data: bytes):
        self._data = data
        self._position = 0

    def read(self, size: int = -1) -> bytes:
        if size < 0:
            size = len(self._data) - self._position
        result = self._data[self._position : self._position + size]
        self._position += len(result)
        return result

    def seek(self, offset: int, whence: int = 0) -> int:
        if whence == 0:
            position = offset
        elif whence == 1:
            position = self._position + offset
        elif whence == 2:
            position = len(self._data) + offset
        else:  # pragma: no cover - standard library uses only valid values
            raise ValueError("invalid whence")
        self._position = max(0, position)
        return self._position

    def tell(self) -> int:
        return self._position

    def seekable(self) -> bool:
        return True


def validate_upload(data: bytes, filename: str, media_type: str) -> tuple[str, str]:
    """Validate size, declared type, extension, and content signature."""
    name = safe_filename(filename)
    if not data:
        raise DocumentRejected("the file is empty")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentRejected("the file exceeds the 10 MB limit")
    if data.startswith((b"MZ", b"\x7fELF")):
        raise DocumentRejected("executable files are not supported")
    suffix = Path(name).suffix.casefold()
    allowed = _ALLOWED_MEDIA_TYPES.get(suffix)
    if allowed is None:
        raise DocumentRejected("only PDF, PPTX, TXT, and Markdown files are supported")
    normalized_media = media_type.split(";", 1)[0].strip().casefold()
    if normalized_media not in allowed:
        raise DocumentRejected("the file type does not match its extension")
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise DocumentRejected("the file is not a valid PDF")
        # Fast rejection before the parser sees common encrypted PDFs. The
        # parser repeats this check authoritatively through ``is_encrypted``.
        if b"/Encrypt" in data:
            raise DocumentRejected("encrypted PDFs are not supported")
    elif suffix == ".pptx":
        if not data.startswith(b"PK"):
            raise DocumentRejected("the file is not a valid PPTX presentation")
        _validate_zip(data)
    else:
        if b"\x00" in data:
            raise DocumentRejected("binary content is not supported as text")
        try:
            data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise DocumentRejected("text files must use UTF-8 encoding") from None
    return name, suffix


def _bounded_chunks(parts: list[tuple[str, str]]) -> list[IntakeDocumentChunk]:
    chunks: list[IntakeDocumentChunk] = []
    remaining = MAX_EXTRACTED_CHARS
    for location, raw in parts:
        if len(chunks) >= MAX_CHUNKS or remaining <= 0:
            break
        cleaned, locally_truncated = ingest(raw)
        if not cleaned:
            continue
        text = cleaned[:remaining]
        aggregate_truncated = len(cleaned) > remaining
        chunks.append(
            IntakeDocumentChunk(
                chunk_id=f"pending:{len(chunks) + 1}",
                location=location,
                text=text,
                truncated=locally_truncated or aggregate_truncated,
            )
        )
        remaining -= len(text)
    if not chunks:
        raise DocumentRejected("no readable text was found; image-only files need OCR")
    return chunks


def _parse_path(path: Path, suffix: str) -> list[IntakeDocumentChunk]:
    if suffix == ".pdf":
        try:
            reader = PdfReader(str(path), strict=True)
            if reader.is_encrypted:
                raise DocumentRejected("encrypted PDFs are not supported")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise DocumentRejected(f"PDFs may contain at most {MAX_PDF_PAGES} pages")
            return _bounded_chunks(
                [(f"page:{index}", page.extract_text() or "") for index, page in enumerate(reader.pages, 1)]
            )
        except DocumentRejected:
            raise
        except Exception:
            raise DocumentRejected("the PDF is malformed or unsupported") from None
    if suffix == ".pptx":
        try:
            deck = Presentation(str(path))
            if len(deck.slides) > MAX_PPTX_SLIDES:
                raise DocumentRejected(
                    f"presentations may contain at most {MAX_PPTX_SLIDES} slides"
                )
            parts: list[tuple[str, str]] = []
            for index, slide in enumerate(deck.slides, 1):
                text = "\n".join(
                    shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text
                )
                parts.append((f"slide:{index}", text))
            return _bounded_chunks(parts)
        except DocumentRejected:
            raise
        except Exception:
            raise DocumentRejected("the presentation is malformed or unsupported") from None
    try:
        text = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeDecodeError):
        raise DocumentRejected("the text file is malformed or unsupported") from None
    return _bounded_chunks([("document", text)])


def extract_upload(
    data: bytes,
    filename: str,
    media_type: str,
    *,
    timeout_seconds: int = PARSER_TIMEOUT_SECONDS,
    temp_directory: Path | None = None,
) -> tuple[str, list[IntakeDocumentChunk]]:
    """Extract in a time-bounded child process and always delete raw bytes."""
    name, suffix = validate_upload(data, filename, media_type)
    raw_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", suffix=suffix, dir=temp_directory, delete=False
        ) as raw:
            raw.write(data)
            raw_path = Path(raw.name)
        command = [sys.executable, "-m", "agent.intake_documents", "--parse", str(raw_path), suffix]
        try:
            result = subprocess.run(
                command,
                check=False,
                capture_output=True,
                text=True,
                timeout=timeout_seconds,
            )
        except subprocess.TimeoutExpired:
            raise DocumentRejected("document extraction timed out") from None
        try:
            payload: dict[str, Any] = json.loads(result.stdout)
        except json.JSONDecodeError:
            raise DocumentRejected("document extraction failed safely") from None
        if result.returncode != 0 or "error" in payload:
            raise DocumentRejected(str(payload.get("error", "document extraction failed safely")))
        chunks = [IntakeDocumentChunk.model_validate(item) for item in payload["chunks"]]
        return name, chunks
    finally:
        if raw_path is not None:
            raw_path.unlink(missing_ok=True)


def _main() -> int:
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--parse", type=Path, required=True)
    parser.add_argument("suffix")
    args = parser.parse_args()
    try:
        # Linux production workers get a hard memory/CPU envelope in addition
        # to the parent's wall-clock timeout. Windows lacks ``resource``;
        # there the disposable child and timeout remain the containment wall.
        try:
            import resource

            memory_limit = 512 * 1024 * 1024
            resource.setrlimit(resource.RLIMIT_AS, (memory_limit, memory_limit))
            resource.setrlimit(resource.RLIMIT_CPU, (PARSER_TIMEOUT_SECONDS, PARSER_TIMEOUT_SECONDS))
        except (ImportError, OSError, ValueError):
            pass
        chunks = _parse_path(args.parse, args.suffix)
        print(json.dumps({"chunks": [chunk.model_dump(mode="json") for chunk in chunks]}))
        return 0
    except DocumentRejected as exc:
        print(json.dumps({"error": str(exc)}))
        return 2


if __name__ == "__main__":  # pragma: no cover - exercised through subprocess
    raise SystemExit(_main())
