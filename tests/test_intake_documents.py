"""Hostile-file defenses and founder ownership for intake evidence."""

from __future__ import annotations

import subprocess
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pptx import Presentation

from agent import config
from agent.intake import new_session
from agent.intake_documents import (
    DocumentRejected,
    MAX_FILE_BYTES,
    extract_upload,
    validate_upload,
)
from api.main import MAX_UPLOAD_BODY_BYTES, app


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("KAIROS_DB_URL", f"sqlite:///{tmp_path}/documents.db")
    monkeypatch.setenv("KAIROS_ALLOW_OPEN_API", "1")
    monkeypatch.delenv("KAIROS_API_TOKEN", raising=False)
    config.settings.cache_clear()
    with TestClient(app) as test_client:
        yield test_client
    config.settings.cache_clear()


def _session(client: TestClient) -> str:
    response = client.post("/founders/founder_demo/intake/sessions")
    assert response.status_code == 200
    return response.json()["session"]["session_id"]


def _pptx_bytes(text: str = "We help university labs coordinate equipment.") -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Kairos"
    slide.placeholders[1].text = text
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("data", "filename", "media_type", "reason"),
    [
        (b"MZ" + b"x" * 20, "brief.txt", "text/plain", "executable"),
        (b"not a pdf", "brief.pdf", "application/pdf", "valid PDF"),
        (b"%PDF-1.7\n/Encrypt", "brief.pdf", "application/pdf", "encrypted"),
        (b"hello", "brief.txt", "application/pdf", "does not match"),
        (b"hello\x00world", "brief.txt", "text/plain", "binary"),
        (b"hello", "brief.exe", "text/plain", "only PDF"),
    ],
)
def test_signature_extension_and_mime_disagreements_are_rejected(
    data, filename, media_type, reason
):
    with pytest.raises(DocumentRejected, match=reason):
        validate_upload(data, filename, media_type)


def test_pptx_macro_payload_is_rejected_before_office_parser_runs():
    source = BytesIO(_pptx_bytes())
    output = BytesIO()
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as changed:
        for info in original.infolist():
            changed.writestr(info, original.read(info.filename))
        changed.writestr("ppt/vbaProject.bin", b"macro")

    with pytest.raises(DocumentRejected, match="macro-enabled"):
        validate_upload(
            output.getvalue(),
            "brief.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )


def test_pptx_decompression_bomb_is_rejected_before_parser_runs():
    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("ppt/presentation.xml", "presentation")
        archive.writestr("ppt/media/payload.bin", b"0" * 1024 * 1024)

    with pytest.raises(DocumentRejected, match="compression ratio"):
        validate_upload(
            output.getvalue(),
            "brief.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )


def test_image_only_pdf_is_reported_as_unsupported_and_deleted(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    output = BytesIO()
    writer.write(output)

    with pytest.raises(DocumentRejected, match="image-only"):
        extract_upload(
            output.getvalue(),
            "scan.pdf",
            "application/pdf",
            temp_directory=tmp_path,
        )
    assert list(tmp_path.iterdir()) == []


def test_pdf_page_count_is_bounded_before_extraction(tmp_path):
    writer = PdfWriter()
    for _ in range(101):
        writer.add_blank_page(width=10, height=10)
    output = BytesIO()
    writer.write(output)

    with pytest.raises(DocumentRejected, match="at most 100 pages"):
        extract_upload(
            output.getvalue(),
            "long.pdf",
            "application/pdf",
            temp_directory=tmp_path,
        )


def test_text_extraction_is_sanitized_bounded_and_raw_temp_is_deleted(tmp_path):
    name, chunks = extract_upload(
        b"<!--hidden--> <script>bad()</script>Real traction: 47 pilots.",
        "../brief.md",
        "text/markdown",
        temp_directory=tmp_path,
    )

    assert name == "brief.md"
    assert chunks[0].location == "document"
    assert chunks[0].text == "Real traction: 47 pilots."
    assert list(tmp_path.iterdir()) == []


def test_pptx_extraction_retains_slide_provenance(tmp_path):
    _, chunks = extract_upload(
        _pptx_bytes("Serving 12 university labs."),
        "deck.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        temp_directory=tmp_path,
    )

    assert chunks[0].location == "slide:1"
    assert "12 university labs" in chunks[0].text
    assert list(tmp_path.iterdir()) == []


def test_parser_timeout_is_safe_and_raw_temp_is_deleted(monkeypatch, tmp_path):
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired(cmd="parser", timeout=1)

    monkeypatch.setattr(subprocess, "run", timeout)
    with pytest.raises(DocumentRejected, match="timed out"):
        extract_upload(
            b"A valid founder brief.",
            "brief.txt",
            "text/plain",
            temp_directory=tmp_path,
        )
    assert list(tmp_path.iterdir()) == []


def test_authenticated_upload_evidence_and_removal_flow(client):
    session_id = _session(client)
    uploaded = client.post(
        f"/founders/founder_demo/intake/sessions/{session_id}/documents",
        files={"file": ("brief.txt", b"We serve 18 campus labs.", "text/plain")},
        data={"role": "admin"},  # ignored; no form field can assign privileges
    )

    assert uploaded.status_code == 200
    document = uploaded.json()
    assert document["status"] == "ready"
    assert document["slot"] == 1
    assert "role" not in document
    chunk_id = document["chunks"][0]["chunk_id"]
    evidence = client.get(
        f"/founders/founder_demo/intake/sessions/{session_id}/evidence/{chunk_id}"
    )
    assert evidence.status_code == 200
    assert evidence.json()["location"] == "document"

    removed = client.delete(
        f"/founders/founder_demo/intake/sessions/{session_id}/documents/{document['document_id']}"
    )
    assert removed.status_code == 204


def test_two_document_limit_is_database_enforced_and_removal_releases_a_slot(client):
    session_id = _session(client)
    ids = []
    for index in range(2):
        response = client.post(
            f"/founders/founder_demo/intake/sessions/{session_id}/documents",
            files={"file": (f"brief-{index}.txt", b"Founder context", "text/plain")},
        )
        assert response.status_code == 200
        ids.append(response.json()["document_id"])

    third = client.post(
        f"/founders/founder_demo/intake/sessions/{session_id}/documents",
        files={"file": ("third.txt", b"More context", "text/plain")},
    )
    assert third.status_code == 409

    client.delete(
        f"/founders/founder_demo/intake/sessions/{session_id}/documents/{ids[0]}"
    )
    retry = client.post(
        f"/founders/founder_demo/intake/sessions/{session_id}/documents",
        files={"file": ("third.txt", b"More context", "text/plain")},
    )
    assert retry.status_code == 200


def test_cross_founder_document_paths_are_indistinguishable_from_missing(client):
    other = new_session("founder_other", None)
    app.state.repo.create_intake_session(other)

    response = client.post(
        f"/founders/founder_other/intake/sessions/{other.session_id}/documents",
        files={"file": ("brief.txt", b"private", "text/plain")},
    )
    assert response.status_code == 404


def test_upload_route_has_a_narrow_request_ceiling(client):
    session_id = _session(client)
    response = client.post(
        f"/founders/founder_demo/intake/sessions/{session_id}/documents",
        headers={"Content-Length": str(MAX_UPLOAD_BODY_BYTES + 1)},
        content=b"x",
    )
    assert response.status_code == 413


def test_file_byte_limit_is_independent_of_multipart_overhead():
    with pytest.raises(DocumentRejected, match="10 MB"):
        validate_upload(b"x" * (MAX_FILE_BYTES + 1), "brief.txt", "text/plain")
